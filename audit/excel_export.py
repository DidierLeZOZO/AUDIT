"""Génération de la fiche Excel « Liste des travaux et entreprises »."""
from __future__ import annotations

import copy
import io
from datetime import date
from pathlib import Path

import openpyxl

from .gabarit import CHEMIN as GABARIT_PAR_DEFAUT
from .gabarit import construire
from .parser import Batiment, Scenario
from .securite import ecrire_cellule_openpyxl

LIGNE_DEBUT = 9      # première ligne de travaux dans le gabarit
LIGNES_DISPO = 6     # lignes prévues dans le gabarit avant insertion


def _copier_style_ligne(ws, src: int, dst: int, col_min: int = 2, col_max: int = 7) -> None:
    for col in range(col_min, col_max + 1):
        s, d = ws.cell(row=src, column=col), ws.cell(row=dst, column=col)
        d.font, d.border, d.fill = copy.copy(s.font), copy.copy(s.border), copy.copy(s.fill)
        d.number_format, d.alignment = s.number_format, copy.copy(s.alignment)


def lignes_travaux(sc: Scenario) -> list[dict]:
    """Aplati les travaux de toutes les étapes du scénario."""
    out = []
    for e in sc.etapes:
        suffix = f" — {e.libelle}" if len(sc.etapes) > 1 else ""
        for t in e.travaux:
            out.append({"nature": f"{t.nature_affichee}{suffix}",
                        "carac": t.caracteristiques, "quantite": t.quantite})
    return out


def generer_excel(batiment: Batiment, sc: Scenario, prefill_realises: bool = False,
                  gabarit: Path | None = None) -> bytes:
    gabarit = gabarit or GABARIT_PAR_DEFAUT
    if not gabarit.exists() and gabarit == GABARIT_PAR_DEFAUT:
        construire(gabarit)
    wb = openpyxl.load_workbook(gabarit)
    ws = wb.active

    ecrire_cellule_openpyxl(ws["C4"], batiment.beneficiaire or "")
    ecrire_cellule_openpyxl(ws["F4"], sc.nom.replace("\n", " "))
    ecrire_cellule_openpyxl(ws["C5"], batiment.adresse or "")
    ws["F5"] = date.today().strftime("%d/%m/%Y")

    travaux = lignes_travaux(sc)
    if len(travaux) > LIGNES_DISPO:
        extra = len(travaux) - LIGNES_DISPO
        ws.insert_rows(LIGNE_DEBUT + LIGNES_DISPO, amount=extra)
        for i in range(extra):
            _copier_style_ligne(ws, LIGNE_DEBUT, LIGNE_DEBUT + LIGNES_DISPO + i)

    for i, t in enumerate(travaux):
        r = LIGNE_DEBUT + i
        ecrire_cellule_openpyxl(ws.cell(row=r, column=2), t["nature"])
        ecrire_cellule_openpyxl(ws.cell(row=r, column=3), t["carac"])
        ecrire_cellule_openpyxl(ws.cell(row=r, column=4), t["quantite"])
        # L'audit ne dit rien des travaux réalisés : pré-remplissage uniquement sur demande.
        if prefill_realises:
            ecrire_cellule_openpyxl(ws.cell(row=r, column=5), t["nature"])
            ecrire_cellule_openpyxl(ws.cell(row=r, column=6), t["carac"])
            ecrire_cellule_openpyxl(ws.cell(row=r, column=7), t["quantite"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
