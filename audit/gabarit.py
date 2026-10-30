"""Génère le gabarit Excel « Liste des travaux et entreprises » par défaut.

Disposition attendue par audit/excel_export.py :
  C4 bénéficiaire · F4 scénario · C5 adresse · F5 date
  ligne 9+ : B nature · C caractéristiques · D quantité (préconisés)
             E nature · F caractéristiques · G quantité (réalisés)

Pour utiliser VOTRE gabarit officiel, remplacez simplement le fichier
audit/donnees/Liste_des_travaux_et_entreprises_Modèle.xlsx en respectant ces cellules.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CHEMIN = Path(__file__).parent / "donnees" / "Liste_des_travaux_et_entreprises_Modèle.xlsx"


def construire(chemin: Path = CHEMIN) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Liste des travaux"
    fin = Side(style="thin", color="999999")
    cadre = Border(left=fin, right=fin, top=fin, bottom=fin)
    gras = Font(bold=True)
    bleu = PatternFill("solid", fgColor="DDEBF7")

    ws["B2"] = "Liste des travaux et entreprises"
    ws["B2"].font = Font(bold=True, size=14)
    for lib, cel in (("Bénéficiaire", "B4"), ("Adresse", "B5")):
        ws[cel] = lib
        ws[cel].font = gras
    for lib, cel in (("Scénario", "E4"), ("Date", "E5")):
        ws[cel] = lib
        ws[cel].font = gras

    ws.merge_cells("B7:D7")
    ws.merge_cells("E7:G7")
    ws["B7"], ws["E7"] = "Travaux préconisés (audit)", "Travaux réalisés"
    for c in ("B7", "E7"):
        ws[c].font, ws[c].fill = gras, bleu
        ws[c].alignment = Alignment(horizontal="center")
    for i, titre in enumerate(["Nature des travaux", "Caractéristiques / Marque et référence",
                               "Surface / Quantité"] * 2):
        c = ws.cell(row=8, column=2 + i, value=titre)
        c.font, c.fill, c.border = gras, bleu, cadre
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in range(9, 15):
        for col in range(2, 8):
            ws.cell(row=r, column=col).border = cadre
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", (2, 45, 38, 16, 45, 38, 16)):
        ws.column_dimensions[col].width = w
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)
    return chemin


if __name__ == "__main__":
    print(construire())
